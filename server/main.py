from fastapi import FastAPI, UploadFile, File, Header
from contextlib import asynccontextmanager
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
import io
import openpyxl

from processData import processExcelFile
from exchangeRateProcess import initialiseDaemon
from cacheManagement import CacheManager

JSON_FILE_PATH = "./exchangeData/exchageRatesCached.json"

@asynccontextmanager
async def lifespan(app: FastAPI):
    print("FastAPI server starting up. Initializing exchange rate daemon and creating cache...")
    Cache = CacheManager()
    initialiseDaemon()
    yield
    print("FastAPI server shutting down...")

app = FastAPI(lifespan=lifespan)

class Item(BaseModel):
    payload_values : list
    
origins = [
    "http://127.0.0.1:5500",
    "http://localhost:5500",
    "http://127.0.0.1:8000",
    "http://localhost:8000",
    "http:/0.0.0.0:8000"
]

app.add_middleware(
    CORSMiddleware,
    allow_origins=origins,      
    allow_credentials=True,
    allow_methods=["*"],        
    allow_headers=["*"],        
)

def flatten_excel_structures(file_bytes):
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes))
    for sheet in wb.worksheets:
        table_names = list(sheet.tables.keys())
        for tbl_name in table_names:
            del sheet.tables[tbl_name]
            
    sanitized_stream = io.BytesIO()
    wb.save(sanitized_stream)
    sanitized_stream.seek(0)
    return sanitized_stream

@app.post("/files")
async def create_file(
    file: list[UploadFile] = File(),
    x_user_currency: str = Header(default="USD", alias="user-Currency")
):
    user_currency = str(x_user_currency).upper().strip()
    processing_result = None
    
    separate_files_response = {}
    for currentFile in file:
        file_extension = distinguishFileENd(currentFile.filename)
        file_bytes = await currentFile.read()
        file_object = flatten_excel_structures(file_bytes)
        
        processing_result = await returnType(file_extension, file_object, user_currency)
        expense_list = processing_result.get("fileObjData", [])

        failed_count = processing_result.get("failedRowCount")
        failed_count = failed_count if isinstance(failed_count, int) else 0
        
        separate_files_response[currentFile.filename]= {
            "file_status": processing_result.get("status", "failed"),
            "failed_rows_count": processing_result.get("failedRowCount", 0),
             "failed_rows_details": processing_result.get("failedRowDetails") if failed_count > 0 else [] ,
            "expense_data": expense_list if isinstance(expense_list, list) else []
        }
    
    return {
        "normalization_base": "EUR",
        "conversion_matrices": processing_result.get("conversionMatrices") if processing_result else {"value":"None"},
        "file_batch_data": separate_files_response
    }

def distinguishFileENd(FilePath):
    return FilePath.split(".")[-1].lower()

async def returnType(file_extension, file_object, user_currency="USD"):
    match file_extension:
        case "xlsx" | "xls":
            execution_report = await processExcelFile(
                file_object, 
                user_currency=user_currency
            )
            if "error" in execution_report:
                return {
                    "file Type": f"excel ({file_extension})", 
                    "status": "error", 
                    "fileObjData": []
                }
                
            return execution_report
        case "pdf":
            return {"file Type": "pdf", "status": "unimplemented"}
        case _:
            return {"file Type": f"unsupported ({file_extension})", "status": "unsupported"}
